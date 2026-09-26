from datetime import date
from pathlib import Path
import importlib.util
import io
import sys
import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent / 'fixtures'))
import master_reference_sample
from backend.app.sheet.parse import ParsedWorkbook, Harvest, Strain
from backend.app.sheet.sinks import build_import_plan
spec = importlib.util.spec_from_file_location('workbook_api', ROOT / 'web/api/workbook.py')
api = importlib.util.module_from_spec(spec)
spec.loader.exec_module(api)


def test_upload_preview_uses_canonical_parser(tmp_path):
    raw = Path(master_reference_sample.build(tmp_path / 'sample.xlsx')).read_bytes()
    plan, preview = api.preview_workbook(raw)
    assert preview['total'] > 10
    assert plan['batches'] and plan['harvests']
    assert all(not row['lot_code'].endswith(('-F1', '-F2')) for row in plan['batches'])
    assert all('_batch_lot_code' in row for row in plan['harvests'])
    assert any(row['notes'] for row in plan['harvests'])
    assert all('batch_id' not in row for row in plan['harvests'])


def test_harvest_import_with_declared_strain_preserves_batch():
    parsed = ParsedWorkbook(strains=[Strain('Oyster', mushroom_type='functional')], harvests=[Harvest('T1-F1', 'T1', 'Oyster', 1, date(2026, 9, 7), 12, 1, 'First flush')])
    plan = build_import_plan(parsed)
    assert plan['strains'][0]['mushroom_type'] == 'functional'
    assert plan['batches'][0]['_ensure_only'] is True
    assert plan['harvests'][0]['_batch_lot_code'] == 'T1'
    assert len(parsed.strains) == 1 and not parsed.batches


@pytest.mark.parametrize('data', [b'', b'not an xlsx', b'x' * (api.MAX_BYTES + 1)])
def test_reject_invalid_file(data):
    with pytest.raises(ValueError): api.preview_workbook(data)


def test_reject_unrecognized_sheet():
    wb = Workbook(); wb.active.append(['Unrelated', 'Data']); wb.active.append(['test', 4])
    raw = io.BytesIO(); wb.save(raw)
    with pytest.raises(ValueError, match='No supported records'): api.preview_workbook(raw.getvalue())


def test_preview_receipt_binds_file_and_prevents_tampering():
    receipt = api.receipt_for('digest', 'secret')
    assert api.verify_receipt(receipt, 'digest', 'secret')
    for digest, secret in [('different-file', 'secret'), ('digest', 'different-key')]:
        with pytest.raises(ValueError): api.verify_receipt(receipt, digest, secret)
    with pytest.raises(ValueError): api.verify_receipt(receipt + '0', 'digest', 'secret')


def test_expired_preview(monkeypatch):
    monkeypatch.setattr(api.time, 'time', lambda: 100)
    receipt = api.receipt_for('digest', 'secret')
    monkeypatch.setattr(api.time, 'time', lambda: 2000)
    with pytest.raises(ValueError, match='expired'): api.verify_receipt(receipt, 'digest', 'secret')


def workbook_bytes(wb):
    raw = io.BytesIO()
    wb.save(raw)
    return raw.getvalue()


@pytest.mark.parametrize('container', [
    'Ordered Sep 17 (Isaac)', 'T-27 (CONFIRMED Sep 15)',
    'T-38 / T-39', '5x5 tub — ASSIGN ID',
])
def test_import_rejects_narrative_and_ambiguous_container_rows(container):
    wb = Workbook(); ws = wb.active; ws.title = 'Grow Cycle Log'
    ws.append(['Strain', 'Tub', 'Flush', 'Inoculated'])
    ws.append(['JAR SUPPLY', container, '', ''])
    with pytest.raises(ValueError, match='Grow Cycle Log row 2'):
        api.preview_workbook(workbook_bytes(wb))


def test_harvest_cannot_bypass_container_validation():
    wb = Workbook(); ws = wb.active; ws.title = 'Harvest Tracker'
    ws.append(['Strain', 'Tub', 'Flush', 'Fresh', 'Harvest Date'])
    ws.append(['Oyster', 'NS block (assign ID)', 1, 30, '2026-09-25'])
    with pytest.raises(ValueError, match='Harvest Tracker row 2'):
        api.preview_workbook(workbook_bytes(wb))


def test_unknown_strain_is_never_created_with_database_default():
    parsed = ParsedWorkbook(harvests=[Harvest('T1-F1', 'T1', 'JAR SUPPLY', 1, date(2026, 9, 7))])
    with pytest.raises(ValueError, match='Unknown strain'):
        build_import_plan(parsed)
    assert parsed.strains == []


def test_one_container_cannot_collapse_different_strains():
    wb = Workbook(); ws = wb.active; ws.title = 'Grow Cycle Log'
    ws.append(['Strain', 'Tub', 'Flush', 'Inoculated'])
    ws.append(['Oyster', 'T-01', 1, '2026-09-01'])
    ws.append(['Enigma', 'T-01', 2, '2026-09-02'])
    with pytest.raises(ValueError, match='conflicting strains'):
        api.preview_workbook(workbook_bytes(wb))


def test_unclassified_library_strain_cannot_default_to_magic():
    wb = Workbook(); ws = wb.active; ws.title = 'Strain Library'
    ws.append(['Strain', 'Status', 'Potency', 'Grow Again', 'Mushroom Type'])
    ws.append(['Reishi', 'Active', '', 'Yes', ''])
    with pytest.raises(ValueError, match='Missing Mushroom Type for Reishi'):
        api.preview_workbook(workbook_bytes(wb))


def test_explicit_functional_library_classification_survives_import():
    wb = Workbook(); ws = wb.active; ws.title = 'Strain Library'
    ws.append(['Strain', 'Status', 'Potency', 'Grow Again', 'Mushroom Type'])
    ws.append(['Reishi', 'Active', '', 'Yes', 'Functional'])
    plan, _ = api.preview_workbook(workbook_bytes(wb))
    assert plan['strains'][0]['mushroom_type'] == 'functional'
    assert plan['strains'][0]['species'] != 'Psilocybe cubensis'
