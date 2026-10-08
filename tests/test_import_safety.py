from datetime import date
import pytest
from openpyxl import Workbook
from backend.app.sheet.parse import ParsedWorkbook, Harvest, parse_harvests
from backend.app.sheet.safety import validate_import, validate_layout


def harvest(**changes):
    values = dict(lot_code='QB-GT-260529-F1', tub='QB-GT-260529', strain='Golden Teacher', flush_number=1, harvested_on=date(2026, 6, 20), fresh_g=67, dry_g=8)
    return Harvest(**(values | changes))


def test_unreadable_and_duplicate_harvests_cannot_import():
    with pytest.raises(ValueError, match='Unreadable'):
        validate_import(ParsedWorkbook(harvests=[harvest(unparsed=['Dry: see jar'])]))
    with pytest.raises(ValueError, match='Duplicate'):
        validate_import(ParsedWorkbook(harvests=[harvest(), harvest()]))
    validate_import(ParsedWorkbook(harvests=[harvest(), harvest(lot_code='pick-2')]))


def test_explicit_harvest_ids_preserve_separate_picks():
    wb = Workbook(); ws = wb.active; ws.title = 'Harvest Tracker'
    ws.append(['Strain', 'Tub', 'Flush', 'Harvest Date', 'Fresh (g)', 'Dry (g)', 'Harvest ID'])
    ws.append(['Golden Teacher', 'QB-GT-260529', 1, '2026-06-20', 67, 8, 'pick-1'])
    ws.append(['Golden Teacher', 'QB-GT-260529', 1, '2026-06-21', 2, 0, 'pick-2'])
    records = parse_harvests(wb)
    assert [r.lot_code for r in records] == ['pick-1', 'pick-2']
    validate_import(ParsedWorkbook(harvests=records))


def test_layout_drift_stops_import():
    wb = Workbook(); wb.active.title = 'Harvest Tracker'; wb.active.append(['Strain', 'Tub', 'Flush', 'Fresh', 'Changed Date Heading'])
    with pytest.raises(ValueError, match='columns'):
        validate_layout(wb)
    wb = Workbook()
    with pytest.raises(ValueError, match='Missing required'):
        validate_layout(wb, master=True)


def test_missing_weight_is_not_explicit_zero():
    wb = Workbook(); ws = wb.active; ws.title = 'Harvest Tracker'
    ws.append(['Strain','Tub','Flush','Harvest Date','Fresh (g)','Dry (g)','Harvest ID'])
    ws.append(['Golden Teacher','QB-GT-260529',1,'2026-06-20',None,8,'pick-1'])
    ws.append(['Golden Teacher','QB-GT-260529',1,'2026-06-21',2,0,'pick-2'])
    records = parse_harvests(wb)
    assert records[0].fresh_recorded is False
    assert records[0].dry_recorded is True
    assert records[1].dry_recorded is True
    validate_import(ParsedWorkbook(harvests=records))
