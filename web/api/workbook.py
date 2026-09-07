"""Direct workbook preview/import using the same parser as the scheduled job."""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
from pathlib import Path
import sys
import time
from http.server import BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs, unquote
from uuid import uuid4
from zipfile import ZipFile, BadZipFile

sys.path.insert(0, str(Path(__file__).parent / '_vendor'))
# Local tests/development can run before npm's postinstall bundles the engine.
sys.path.append(str(Path(__file__).resolve().parents[2]))
from openpyxl import load_workbook
from backend.app.sheet.parse import parse_workbook
from backend.app.sheet.sinks import build_import_plan, SupabaseSink

MAX_BYTES = 4 * 1024 * 1024


def preview_workbook(data: bytes) -> tuple[dict, dict]:
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Choose an .xlsx workbook up to 4 MB.')
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(f.file_size for f in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('This workbook expands beyond the 40 MB safety limit. Split it into smaller workbooks.')
            if '[Content_Types].xml' not in archive.namelist():
                raise ValueError('This is not an Excel workbook. Save it as .xlsx and try again.')
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('This file could not be opened. Save an unencrypted .xlsx copy and try again.') from exc
    try:
        for ws in wb.worksheets:
            if (ws.max_row or 0) > 20000 or (ws.max_column or 0) > 200 or (ws.max_row or 0) * (ws.max_column or 0) > 500000:
                raise ValueError('A sheet is too large. Remove unused rows and columns, then try again.')
        parsed = parse_workbook(wb)
        plan = build_import_plan(parsed)
        counts = {table: len(rows) for table, rows in plan.items() if rows}
        if not counts:
            raise ValueError('No supported records found. Use the Master Cultivation Reference layout; keep its tab names and column headings.')
        skipped = sum(1 for h in parsed.harvests if not h.harvested_on)
        warnings = [f'{skipped} harvest row(s) have no date and will be skipped.'] if skipped else []
        return plan, {
            'counts': counts, 'total': sum(counts.values()), 'sheets': wb.sheetnames,
            'warnings': warnings, 'sha256': hashlib.sha256(data).hexdigest(),
        }
    finally:
        wb.close()


def receipt_for(digest: str, key: str) -> str:
    body = f'{digest}:{int(time.time()) + 1800}:{uuid4()}'
    signature = hmac.new(key.encode(), body.encode(), hashlib.sha256).hexdigest()
    return f'{body}:{signature}'


def verify_receipt(receipt: str, digest: str, key: str) -> str:
    try:
        actual, expiry, request_id, signature = receipt.split(':')
        body = f'{actual}:{expiry}:{request_id}'
        expected = hmac.new(key.encode(), body.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected) or actual != digest or int(expiry) < time.time():
            raise ValueError()
        return request_id
    except (ValueError, TypeError):
        raise ValueError('The file changed or its preview expired. Preview it again before importing.')


class handler(BaseHTTPRequestHandler):
    def respond(self, status: int, payload: dict):
        body = json.dumps(payload).encode()
        self.send_response(status)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        origin = self.headers.get('Origin', '')
        host = self.headers.get('X-Forwarded-Host') or self.headers.get('Host', '')
        if not origin or urlparse(origin).netloc != host:
            return self.respond(403, {'ok': False, 'message': 'Cross-origin import is not allowed.'})
        try:
            length = int(self.headers.get('Content-Length', '0'))
        except ValueError:
            length = 0
        if length <= 0 or length > MAX_BYTES:
            return self.respond(413, {'ok': False, 'message': 'Choose an .xlsx workbook up to 4 MB.'})
        name = unquote(self.headers.get('X-Workbook-Name', 'workbook.xlsx'))[:255]
        if not name.lower().endswith('.xlsx'):
            return self.respond(422, {'ok': False, 'message': 'Save this spreadsheet as .xlsx before uploading.'})
        key = os.environ.get('SUPABASE_SERVICE_ROLE_KEY', '')
        url = os.environ.get('NEXT_PUBLIC_SUPABASE_URL') or os.environ.get('SUPABASE_URL', '')
        if not key or not url:
            return self.respond(503, {'ok': False, 'message': 'Workbook import is temporarily unavailable. Please contact morphiclabsdata@gmail.com.'})
        try:
            data = self.rfile.read(length)
            plan, preview = preview_workbook(data)
            mode = parse_qs(urlparse(self.path).query).get('mode', ['preview'])[0]
            if mode == 'preview':
                return self.respond(200, {'ok': True, **preview, 'receipt': receipt_for(preview['sha256'], key)})
            if mode != 'import':
                raise ValueError('Invalid import request.')
            request_id = verify_receipt(self.headers.get('X-Import-Receipt', ''), preview['sha256'], key)
            with SupabaseSink(url, key, timeout=45) as sink:
                response = sink.client.post(f'{sink.base}/rpc/import_workbook', json={
                    'p_tables': plan, 'p_source': f'Upload: {name}', 'p_request_id': request_id,
                })
                if not response.is_success:
                    # Avoid logging workbook contents, customer details, or credentials.
                    print(f'workbook import RPC failed: HTTP {response.status_code}', file=sys.stderr)
                    return self.respond(422, {'ok': False, 'message': 'The workbook could not be saved. No records were changed. Check required values and try again, or contact morphiclabsdata@gmail.com.'})
                counts = response.json()
            return self.respond(200, {'ok': True, 'counts': counts, 'total': sum(counts.values()), 'message': 'Workbook imported. Your records are ready.'})
        except ValueError as exc:
            return self.respond(422, {'ok': False, 'message': str(exc)})
        except Exception as exc:
            print(f'workbook import failed: {type(exc).__name__}', file=sys.stderr)
            return self.respond(503, {'ok': False, 'message': 'Import was interrupted. Your file is still selected; retry to safely recover the result.'})


if __name__ == '__main__':
    from http.server import ThreadingHTTPServer
    ThreadingHTTPServer(('127.0.0.1', int(os.environ.get('WORKBOOK_PORT', '3101'))), handler).serve_forever()
